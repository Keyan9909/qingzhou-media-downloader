"""Create the polished Word edition of the beginner user guide."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "详细使用教程.md"
OUTPUT = ROOT / "release" / "轻舟下载器_1.0.6" / "轻舟下载器_1.0.6_详细使用教程.docx"

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
NAVY = "203748"
MUTED = "667085"
LIGHT_BLUE = "E8EEF5"
LIGHT_GOLD = "FFF4E5"
GOLD = "9A6700"
WHITE = "FFFFFF"
BLACK = "111827"
ASCII_FONT = "Calibri"
CJK_FONT = "Microsoft YaHei"


def set_run_font(run, *, size=None, bold=None, color=None, italic=None, ascii_font=ASCII_FONT):
    run.font.name = ascii_font
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), CJK_FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), ascii_font)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), ascii_font)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, *, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        item = tc_mar.find(qn(f"w:{edge}"))
        if item is None:
            item = OxmlElement(f"w:{edge}")
            tc_mar.append(item)
        item.set(qn("w:w"), str(value))
        item.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths, *, indent=120):
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent))
    tbl_ind.set(qn("w:type"), "dxa")
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            cell.width = Inches(widths[index] / 1440)
            tc_w = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
            tc_w.set(qn("w:w"), str(widths[index]))
            tc_w.set(qn("w:type"), "dxa")
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)


def add_page_field(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])
    set_run_font(run, size=9, color=MUTED)


def set_keep_with_next(paragraph):
    paragraph.paragraph_format.keep_with_next = True


def add_inline_markdown(paragraph, text, *, size=11, color=BLACK, bold=False):
    parts = re.split(r"(`[^`]+`)", text)
    for part in parts:
        if not part:
            continue
        is_code = part.startswith("`") and part.endswith("`")
        value = part[1:-1] if is_code else part
        run = paragraph.add_run(value)
        set_run_font(
            run,
            size=9.5 if is_code else size,
            color=DARK_BLUE if is_code else color,
            bold=bold,
            ascii_font="Consolas" if is_code else ASCII_FONT,
        )


def add_numbering_definition(doc, kind):
    numbering = doc.part.numbering_part.element
    abstract_ids = [int(item.get(qn("w:abstractNumId"))) for item in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(item.get(qn("w:numId"))) for item in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids, default=-1) + 1
    num_id = max(num_ids, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    level.append(start)
    fmt = OxmlElement("w:numFmt")
    fmt.set(qn("w:val"), "decimal" if kind == "number" else "bullet")
    level.append(fmt)
    text = OxmlElement("w:lvlText")
    text.set(qn("w:val"), "%1." if kind == "number" else "●")
    level.append(text)
    jc = OxmlElement("w:lvlJc")
    jc.set(qn("w:val"), "left")
    level.append(jc)
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "540")
    tabs.append(tab)
    p_pr.append(tabs)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "540")
    ind.set(qn("w:hanging"), "270")
    p_pr.append(ind)
    level.append(p_pr)
    abstract.append(level)
    numbering.append(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def add_list_item(doc, text, num_id):
    paragraph = doc.add_paragraph(style="Normal")
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), str(num_id))
    num_pr.extend([ilvl, num])
    p_pr.append(num_pr)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.line_spacing = 1.25
    paragraph.paragraph_format.keep_together = True
    add_inline_markdown(paragraph, text)
    return paragraph


def add_callout(doc, label, text, *, fill=LIGHT_BLUE, accent=BLUE):
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    set_table_geometry(table, [9360])
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.15
    run = paragraph.add_run(f"{label}  ")
    set_run_font(run, size=11, bold=True, color=accent)
    add_inline_markdown(paragraph, text, size=10.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = ASCII_FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT)
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(BLACK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, DARK_BLUE, 10, 5),
    ):
        style = doc.styles[name]
        style.font.name = ASCII_FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    header = section.header
    header_p = header.paragraphs[0]
    header_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    header_p.paragraph_format.space_after = Pt(0)
    run = header_p.add_run("轻舟下载器  |  完整使用教程")
    set_run_font(run, size=9, color=MUTED, bold=True)

    footer = section.footer
    footer_p = footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer_p.paragraph_format.space_after = Pt(0)
    prefix = footer_p.add_run("轻舟下载器 1.0.6  ·  第 ")
    set_run_font(prefix, size=9, color=MUTED)
    add_page_field(footer_p)
    suffix = footer_p.add_run(" 页")
    set_run_font(suffix, size=9, color=MUTED)


def add_cover(doc):
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(105)

    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    kicker.paragraph_format.space_after = Pt(16)
    set_run_font(kicker.add_run("全平台内容保存工具"), size=11, bold=True, color=BLUE)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(10)
    set_run_font(title.add_run("轻舟下载器"), size=30, bold=True, color=NAVY)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(8)
    set_run_font(subtitle.add_run("完整使用教程（小白版）"), size=18, bold=True, color=DARK_BLUE)

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.paragraph_format.space_after = Pt(46)
    set_run_font(meta.add_run("适用于 1.0.6 及以上版本  |  Windows"), size=10.5, color=MUTED)

    lead = doc.add_paragraph()
    lead.alignment = WD_ALIGN_PARAGRAPH.CENTER
    lead.paragraph_format.space_after = Pt(14)
    set_run_font(lead.add_run("从安装、登录到下载和故障处理，按步骤操作即可。"), size=12, color=BLACK)

    warning = doc.add_paragraph()
    warning.alignment = WD_ALIGN_PARAGRAPH.CENTER
    warning.paragraph_format.space_after = Pt(75)
    set_run_font(warning.add_run("YouTube 用户请提前安装并登录 Firefox 浏览器"), size=12, bold=True, color=GOLD)

    legal = doc.add_paragraph()
    legal.alignment = WD_ALIGN_PARAGRAPH.CENTER
    legal.paragraph_format.space_after = Pt(0)
    set_run_font(
        legal.add_run("仅用于你拥有下载、备份或其他合法使用权限的内容"),
        size=9.5,
        color=MUTED,
    )
    doc.add_page_break()


def add_quick_start(doc):
    heading = doc.add_paragraph("开始前先看这里", style="Heading 1")
    heading.paragraph_format.space_before = Pt(0)
    add_callout(
        doc,
        "YouTube 必看",
        "先在 Firefox 登录 YouTube，确认目标视频能播放，然后完全关闭 Firefox，再回到轻舟下载器开始下载。软件内置浏览器不用于 YouTube 登录。",
        fill=LIGHT_GOLD,
        accent=GOLD,
    )

    doc.add_paragraph("各平台登录方式速查", style="Heading 2")
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    set_table_geometry(table, [2100, 7260])
    headers = table.rows[0].cells
    headers[0].text = "平台"
    headers[1].text = "登录与下载入口"
    for cell in headers:
        set_cell_shading(cell, LIGHT_BLUE)
        for run in cell.paragraphs[0].runs:
            set_run_font(run, size=10.5, bold=True, color=DARK_BLUE)
    rows = [
        ("YouTube", "Firefox 登录并确认视频可播放，完全退出 Firefox 后下载"),
        ("哔哩哔哩", "公开视频可直接下载；受限内容可读取 Firefox、Edge 或 Chrome 登录状态"),
        ("抖音", "使用软件弹出的内置浏览器完成登录或验证"),
        ("小红书", "使用内置浏览器；图文笔记会保存图片、正文和作品信息"),
        ("小鹅通", "使用内置浏览器，以购买课程时相同的身份进入“我的课程”"),
    ]
    for platform, method in rows:
        cells = table.add_row().cells
        cells[0].text = platform
        cells[1].text = method
        for index, cell in enumerate(cells):
            for run in cell.paragraphs[0].runs:
                set_run_font(run, size=10, bold=index == 0, color=BLACK)
    set_table_geometry(table, [2100, 7260])
    doc.add_paragraph()

    doc.add_paragraph("阅读导航", style="Heading 2")
    topics = [
        "首次打开与保存设置",
        "下载中心每个按钮的含义",
        "YouTube 与 Firefox 的完整操作",
        "B站、抖音、小红书视频与图文",
        "小鹅通已购课程",
        "进度、取消、任务记录与 AI/CLI",
        "常见问题和技术支持信息",
    ]
    bullet_id = add_numbering_definition(doc, "bullet")
    for topic in topics:
        add_list_item(doc, topic, bullet_id)
    doc.add_page_break()


def add_source_content(doc):
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    current_kind = None
    current_num_id = None

    for raw in lines[1:]:
        line = raw.strip()
        if not line:
            current_kind = None
            current_num_id = None
            continue
        if line.startswith("### "):
            doc.add_paragraph(line[4:], style="Heading 2")
            current_kind = None
            continue
        if line.startswith("## "):
            heading_text = line[3:]
            paragraph = doc.add_paragraph(heading_text, style="Heading 1")
            if any(key in heading_text for key in ("YouTube 下载", "小鹅通已购", "让 AI")):
                paragraph.paragraph_format.page_break_before = True
            current_kind = None
            continue
        numbered = re.match(r"^\d+\.\s+(.*)$", line)
        bullet = re.match(r"^-\s+(.*)$", line)
        if numbered:
            if current_kind != "number":
                current_num_id = add_numbering_definition(doc, "number")
                current_kind = "number"
            add_list_item(doc, numbered.group(1), current_num_id)
            continue
        if bullet:
            if current_kind != "bullet":
                current_num_id = add_numbering_definition(doc, "bullet")
                current_kind = "bullet"
            add_list_item(doc, bullet.group(1), current_num_id)
            continue

        paragraph = doc.add_paragraph(style="Normal")
        paragraph.paragraph_format.keep_together = True
        add_inline_markdown(paragraph, line)
        current_kind = None
        current_num_id = None

    add_callout(
        doc,
        "隐私提醒",
        "反馈问题时绝对不要发送密码、Cookie、短信验证码、二维码登录截图、私钥或银行卡信息。",
        fill=LIGHT_GOLD,
        accent=GOLD,
    )


def set_core_properties(doc):
    props = doc.core_properties
    props.title = "轻舟下载器完整使用教程（小白版）"
    props.subject = "轻舟下载器 1.0.6 使用说明"
    props.author = "轻舟下载器"
    props.keywords = "轻舟下载器, YouTube, Firefox, 抖音, 小红书, 哔哩哔哩, 小鹅通"
    props.comments = "面向初次使用者的分步骤操作手册"


def main():
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_document(doc)
    set_core_properties(doc)
    add_cover(doc)
    add_quick_start(doc)
    add_source_content(doc)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
